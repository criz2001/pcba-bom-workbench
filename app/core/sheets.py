"""表格读取：统一把 .xlsx / .csv 读成行列表。"""

from __future__ import annotations

import csv
import io
import re
from typing import Any

MAX_ROWS = 200000


class SheetError(RuntimeError):
    pass


def load_rows(path: str, sheet: int | str | None = None) -> tuple[list[list[Any]], list[str]]:
    """读取表格，返回 ``(rows, sheet_names)``。

    第一行即表头，调用方自行决定如何映射列。
    """
    lower = path.lower()
    if lower.endswith(".csv") or lower.endswith(".tsv"):
        return _load_delimited(path), ["CSV"]
    if lower.endswith(".xls"):
        raise SheetError("不支持旧版 .xls，请另存为 .xlsx 后重试")
    if not lower.endswith(".xlsx") and not lower.endswith(".xlsm"):
        raise SheetError("仅支持 .xlsx / .csv 表格文件")
    return _load_xlsx(path, sheet)


def _load_delimited(path: str) -> list[list[Any]]:
    raw = open(path, "rb").read()
    for enc in ("utf-8-sig", "gbk", "utf-8", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover
        text = raw.decode("utf-8", "replace")
    sample = text[:4096]
    dialect = csv.Sniffer().sniff(sample, delimiters=",\t;") if sample else csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    return [list(r) for r in reader][:MAX_ROWS]


def _load_xlsx(path: str, sheet: int | str | None) -> tuple[list[list[Any]], list[str]]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        names = list(wb.sheetnames)
        if sheet is None:
            ws = wb.worksheets[0]
        elif isinstance(sheet, int):
            ws = wb.worksheets[sheet]
        else:
            if sheet not in wb.sheetnames:
                raise SheetError(f"工作簿中不存在工作表「{sheet}」")
            ws = wb[sheet]
        rows: list[list[Any]] = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= MAX_ROWS:
                break
            rows.append(list(row))
        return rows, names
    finally:
        wb.close()


def cell(rows: list[list[Any]], r: int, c: int) -> Any:
    """1 基行列取值，越界返回 None。"""
    if r < 1 or r > len(rows):
        return None
    row = rows[r - 1]
    if c < 1 or c > len(row):
        return None
    return row[c - 1]


def text_of(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def find_header_row(rows: list[list[Any]], keywords: list[str], scan: int = 12) -> int:
    """在前 ``scan`` 行里找出含最多关键字的行，返回 1 基行号（找不到返回 1）。"""
    best_row, best_hits = 1, 0
    for i, row in enumerate(rows[:scan], start=1):
        joined = " ".join(text_of(v) for v in row)
        hits = sum(1 for k in keywords if k in joined)
        if hits > best_hits:
            best_row, best_hits = i, hits
    return best_row


def match_column(header: list[str], *candidates: str, exclude: tuple[str, ...] = ()) -> int:
    """按关键字匹配表头列，返回 1 基列号（0 表示未匹配）。

    先精确匹配，再包含匹配；``exclude`` 中的关键字命中的列会被跳过。
    """
    norm = [re.sub(r"[\s（）()：:/_\-]+", "", str(h or "")).upper() for h in header]
    for cand in candidates:
        c = re.sub(r"[\s（）()：:/_\-]+", "", cand).upper()
        for i, h in enumerate(norm, start=1):
            if h == c and not any(e.upper() in h for e in exclude):
                return i
    for cand in candidates:
        c = re.sub(r"[\s（）()：:/_\-]+", "", cand).upper()
        for i, h in enumerate(norm, start=1):
            if c and c in h and not any(e.upper() in h for e in exclude):
                return i
    return 0
