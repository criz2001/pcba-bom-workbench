"""BOM 模板读取与生成结果写回。

模板只决定列结构、字段名称、格式与固定值；不沿用模板中的物料内容、位号、数量、
BOM 版本与父项信息。
"""

from __future__ import annotations

import os
from typing import Any

from .sheets import find_header_row, load_rows, match_column, text_of

TEMPLATE_HEADER_KEYS = ["项次", "子项物料编码", "子项物料名称", "位置号", "用量"]

FIELD_MATCHERS: list[tuple[str, tuple[str, ...]]] = [
    ("item", ("项次", "序号")),
    ("code", ("子项物料编码", "物料编码", "子项编码")),
    ("name", ("子项物料名称", "物料名称", "子项名称")),
    ("spec", ("子项规格型号", "规格型号", "规格")),
    ("unit", ("子项单位", "基本单位", "单位")),
    ("qty", ("用量:分子", "用量:分子", "用量分子", "用量", "数量")),
    ("denom", ("用量:分母", "用量分母", "分母")),
    ("refs", ("位置号", "位号")),
    ("note", ("备注", "说明")),
    ("parentCode", ("父项物料编码", "父项编码")),
    ("bomVersion", ("BOM版本", "版本")),
    ("bomShortName", ("BOM简称", "简称")),
    ("status", ("数据状态", "状态")),
]

DEFAULT_TEMPLATE_NAME = "标准 BOM 模板（内置）"


class BomTemplate:
    def __init__(self, path: str | None, sheet: str, header_row: int, header: list[str],
                 fields: dict[str, int], meta: dict[str, Any], display_name: str | None = None):
        self.path = path
        self.sheet = sheet
        self.header_row = header_row
        self.header = header
        self.fields = fields
        self.meta = meta
        self.display_name = display_name or (
            os.path.basename(path) if path else DEFAULT_TEMPLATE_NAME
        )

    @property
    def name(self) -> str:
        return self.display_name

    @property
    def column_count(self) -> int:
        return len([h for h in self.header if text_of(h)])

    def summary(self) -> dict[str, Any]:
        return {
            "name": self.name,
            "path": self.path,
            "sheet": self.sheet,
            "headerRow": self.header_row,
            "columns": [text_of(h) for h in self.header],
            "columnCount": self.column_count,
            "fields": self.fields,
            **self.meta,
        }


def _fields_from_header(header: list[str]) -> dict[str, int]:
    fields: dict[str, int] = {}
    for key, candidates in FIELD_MATCHERS:
        idx = match_column(header, *candidates)
        if idx:
            fields[key] = idx
    return fields


def inspect_template(path: str, sheet: int | str | None = None,
                     display_name: str | None = None) -> BomTemplate:
    rows, sheet_names = load_rows(path, sheet)
    if not rows:
        raise ValueError("模板文件为空")
    hdr = find_header_row(rows, TEMPLATE_HEADER_KEYS)
    header = [text_of(v) for v in rows[hdr - 1]]

    meta: dict[str, Any] = {}
    try:
        import openpyxl

        wb = openpyxl.load_workbook(path)
        ws = wb[sheet_names[0]] if sheet is None else (wb[sheet] if isinstance(sheet, str) else wb.worksheets[sheet])
        meta["freeze"] = bool(ws.freeze_panes)
        meta["filter"] = bool(ws.auto_filter and ws.auto_filter.ref)
        wb.close()
    except Exception:
        meta["freeze"] = False
        meta["filter"] = False

    meta["sheets"] = sheet_names
    sheet_name = sheet_names[0] if isinstance(sheet, (int, type(None))) else str(sheet)

    return BomTemplate(path, sheet_name, hdr, header, _fields_from_header(header), meta,
                       display_name)


def builtin_template() -> BomTemplate:
    header = ["项次", "子项物料编码", "子项物料名称", "子项规格型号", "子项单位",
              "用量:分子", "用量:分母", "位置号", "备注"]
    return BomTemplate(None, "BOM", 1, header, _fields_from_header(header),
                       {"freeze": True, "filter": True, "sheets": ["BOM"], "builtin": True})


def _safe_cell(ws, row: int, col: int):
    """取单元格；落在合并区域内的坐标返回 None。

    真实公司模板的表头常带合并单元格，直接对 MergedCell 赋值/设样式
    openpyxl 会抛只读异常——一行样式问题不该让整份 BOM 生成失败。
    """
    from openpyxl.cell.cell import MergedCell

    try:
        cell = ws.cell(row, col)
    except Exception:
        return None
    return None if isinstance(cell, MergedCell) else cell


def _set(cell, **style) -> None:
    """逐项设置样式，任何一项失败都跳过，不影响其它设置。"""
    if cell is None:
        return
    for name, value in style.items():
        try:
            setattr(cell, name, value)
        except Exception:
            pass


def write_bom(template: BomTemplate, out_path: str, rows: list[dict[str, Any]],
              header_meta: dict[str, str] | None = None) -> str:
    """把生成的 BOM 行写入模板副本，保持模板的列顺序与基本视觉风格。"""
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    meta = header_meta or {}

    if template.path:
        wb = openpyxl.load_workbook(template.path)
        ws = wb[template.sheet] if template.sheet in wb.sheetnames else wb.worksheets[0]
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = template.sheet

    hdr = template.header_row
    fields = template.fields
    thin = Side(style="thin", color="BFBFBF")
    border = Border(thin, thin, thin, thin)

    if not template.path:
        style_head = Font(bold=True, color="FFFFFF")
        fill = PatternFill("solid", fgColor="4472C4")
        for j, name in enumerate(template.header, 1):
            c = _safe_cell(ws, hdr, j)
            if c is None:
                continue
            c.value = name
            _set(c, font=style_head, fill=fill, border=border,
                 alignment=Alignment(vertical="center"))
            try:
                ws.column_dimensions[get_column_letter(j)].width = max(
                    12, min(46, len(name) * 2 + 10))
            except Exception:
                pass
        try:
            ws.freeze_panes = ws.cell(hdr + 1, 1).coordinate
            ws.auto_filter.ref = f"A{hdr}:{get_column_letter(len(template.header))}{hdr}"
        except Exception:
            pass

    # 记录模板首行数据行的样式，用于回填
    sample_styles: dict[int, Any] = {}
    for j in range(1, len(template.header) + 1):
        cell = _safe_cell(ws, hdr + 1, j)
        if cell is not None and cell.has_style:
            sample_styles[j] = cell._style

    warn_fill = PatternFill("solid", fgColor="FFEB9C")
    warn_font = Font(color="9C6500")

    # 清空模板里可能存在的示例数据行（模板内容不得沿用）
    max_row = ws.max_row or hdr
    if template.path and max_row > hdr:
        try:
            ws.delete_rows(hdr + 1, max_row - hdr)
        except Exception:
            pass

    def put(row: int, key: str, value: Any) -> None:
        idx = fields.get(key)
        if not idx:
            return
        c = _safe_cell(ws, row, idx)
        if c is None:
            return
        try:
            c.value = value
        except Exception:
            return
        if sample_styles.get(idx) is not None:
            try:
                c._style = sample_styles[idx]
            except Exception:
                pass
        _set(c, border=border)

    for i, data in enumerate(rows):
        r = hdr + 1 + i
        for key in ("item", "code", "name", "spec", "unit", "qty", "denom", "refs", "note"):
            put(r, key, data.get(key, ""))
        if not data.get("pending"):
            continue
        # 未匹配到编码的行：留空编码（避免被误导入），但用底色 + 批注 + 备注三处标注
        for key in ("code", "name", "spec", "note"):
            idx = fields.get(key)
            if not idx:
                continue
            _set(_safe_cell(ws, r, idx), fill=warn_fill, font=warn_font, border=border)
        if data.get("unmatched") and fields.get("code"):
            cell = _safe_cell(ws, r, fields["code"])
            if cell is not None:
                try:
                    cell.comment = Comment(
                        "未匹配到物料编码，请人工指定。\n"
                        + str(data.get("note", "")).replace("；", "\n"),
                        "PCBA BOM 工作台",
                    )
                except Exception:
                    pass

    # 父项 / 版本等模板固定位：仅在调用方显式提供时写入
    if meta:
        for key in ("parentCode", "bomVersion", "bomShortName", "status"):
            value = meta.get(key)
            if not value or not fields.get(key):
                continue
            c = _safe_cell(ws, hdr, fields[key])
            if c is not None:
                try:
                    c.value = value
                except Exception:
                    pass
            _set(c, border=border)

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)
    return out_path


def write_bom_with_fallback(template: BomTemplate, out_path: str, rows: list[dict[str, Any]],
                            header_meta: dict[str, str] | None = None) -> str:
    """优先按模板输出；模板不可用时退回内置模板，保证一定产出 BOM。

    返回警告文案（空字符串表示套用模板成功）。模板损坏、是改了扩展名的 .xls、
    被加密、结构异常等，都不应该让「BOM 无法生成」。
    """
    try:
        write_bom(template, out_path, rows, header_meta)
        return ""
    except Exception as exc:
        if template.path is None:
            raise
        warning = f"模板「{template.name}」无法套用（{exc}），已改用内置 9 列模板输出"
        try:
            os.remove(out_path)
        except OSError:
            pass
        write_bom(builtin_template(), out_path, rows, header_meta)
        return warning
