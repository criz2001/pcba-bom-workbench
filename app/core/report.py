"""核对结果 Excel 输出（4 个 Sheet，含问题高亮）。"""

from __future__ import annotations

import os
from typing import Any

from .checker import CheckResult, parse_cols, problem_rows

RED = "FFC7CE"
YEL = "FFEB9C"
BLU = "DDEBF7"
HEAD = "4472C4"
BORDER = "BFBFBF"

LEVEL_FILL = {"严重": RED, "一般": YEL, "提示": BLU}
LEVEL_FONT = {"严重": "9C0006", "一般": "9C6500", "提示": "1F4E79"}

# 问题类型 → 涉及的 BOM 列（用于高亮）
KIND_COLS = {
    "封装不符": ["spec", "name"],
    "名称与总表不符": ["name"],
    "缺物料编码": ["code"],
    "编码不存在": ["code"],
    "物料已禁用": ["code"],
    "用量≠位号数": ["qty", "refs"],
    "位号格式": ["refs"],
    "位号不存在": ["refs"],
    "位号重复": ["refs"],
    "空行(未填写)": ["code", "name", "refs"],
    "缺项次": ["item"],
}


def write_report(bom_path: str, out_path: str, result: CheckResult, bom,
                 options: dict[str, Any] | None = None) -> str:
    opts = options or {}
    outputs = opts.get("outputs") or {}
    want_highlight = outputs.get("highlight", True)
    want_problems = outputs.get("problemSheet", True)
    want_refs = outputs.get("refSheet", True)
    want_values = outputs.get("valueSheet", True)

    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    thin = Side(style="thin", color=BORDER)
    border = Border(thin, thin, thin, thin)
    red_fill = PatternFill("solid", fgColor=RED)
    yel_fill = PatternFill("solid", fgColor=YEL)
    blu_fill = PatternFill("solid", fgColor=BLU)
    red_font = Font(color=LEVEL_FONT["严重"], bold=True)
    yel_font = Font(color=LEVEL_FONT["一般"])
    blue_font = Font(color=LEVEL_FONT["提示"])

    wb = openpyxl.load_workbook(bom_path)
    ws = wb.worksheets[0]

    # ---------- 1. 原 BOM 表：高亮 + 核对结果列 ----------
    if want_highlight:
        cols = bom.columns
        result_col = max(
            [c for c in cols.values() if c] + [ws.max_column or 1]
        ) + 2
        head = ws.cell(1, result_col, "核对结果（自动核对）")
        head.font = Font(bold=True, color="FFFFFF")
        head.fill = PatternFill("solid", fgColor=HEAD)
        head.alignment = Alignment(vertical="center", wrap_text=True)

        by_row = problem_rows(result)
        # 位置号取自备注列的行：补一条提示，让人知道核对依据不是位置号列
        for item in getattr(result, "refs_from_note", []):
            by_row.setdefault(item["row"], []).append(
                {
                    "kind": "位号取自备注",
                    "level": "提示",
                    "cols": "M",
                    "msg": "位置号列为空，已从备注取回位号 "
                           + " ".join(item["refs"]) + " 并据此核对（不计入位号缺失）",
                }
            )
        for row, problems in by_row.items():
            worst = "严重" if any(p["level"] == "严重" for p in problems) else (
                "一般" if any(p["level"] == "一般" for p in problems) else "提示")
            cell = ws.cell(row, result_col)
            cell.value = "\n".join(f"• [{p['kind']}] {p['msg']}" for p in problems)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.fill = {"严重": red_fill, "一般": yel_fill, "提示": blu_fill}[worst]
            cell.font = {"严重": red_font, "一般": yel_font, "提示": blue_font}[worst]
            cell.border = border

            for p in problems:
                keys = KIND_COLS.get(p["kind"], [])
                if not keys:
                    keys = [{"E": "item", "F": "code", "G": "name",
                             "H": "spec", "J": "qty", "M": "refs"}.get(c.lower(), "")
                            for c in parse_cols(p["cols"])]
                for key in keys:
                    idx = cols.get(key)
                    if not idx:
                        continue
                    target = ws.cell(row, idx)
                    target.fill = {"严重": red_fill, "一般": yel_fill, "提示": blu_fill}[p["level"]]
                    target.font = {"严重": red_font, "一般": yel_font, "提示": blue_font}[p["level"]]

        ws.column_dimensions[get_column_letter(result_col)].width = 62

    # ---------- 2. 问题清单 ----------
    if want_problems:
        ws2 = wb.create_sheet("问题清单")
        headings = ["序号", "级别", "问题类型", "BOM行号", "涉及列", "问题描述", "处理建议"]
        for j, h in enumerate(headings, 1):
            c = ws2.cell(1, j, h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=HEAD)
            c.border = border
        for i, p in enumerate(result.problems, 1):
            values = [i, p["level"], p["kind"],
                      f"第{p['row']}行" if p["row"] else "全局",
                      p["cols"], p["msg"], p["advice"]]
            for j, v in enumerate(values, 1):
                c = ws2.cell(i + 1, j, v)
                c.alignment = Alignment(wrap_text=True, vertical="top")
                c.border = border
                if p["level"] == "严重":
                    c.fill, c.font = red_fill, Font(color=LEVEL_FONT["严重"])
                elif p["level"] == "一般":
                    c.fill, c.font = yel_fill, yel_font
        for col, width in zip("ABCDEFG", [5, 6, 14, 9, 9, 62, 46]):
            ws2.column_dimensions[col].width = width
        ws2.freeze_panes = "A2"

    # ---------- 3. 位号对照 ----------
    if want_refs:
        ws3 = wb.create_sheet("位号对照")
        for j, h in enumerate(["位号", "PCB封装(实际)", "原理图标注值", "BOM是否包含", "核对结论"], 1):
            c = ws3.cell(1, j, h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=HEAD)
            c.border = border
        for i, item in enumerate(result.ref_table, 2):
            values = [item["ref"], item["pcbFootprint"], item["schValue"],
                      "是" if item["inBom"] else "否", item["conclusion"]]
            for j, v in enumerate(values, 1):
                c = ws3.cell(i, j, v)
                c.border = border
                if not item["inBom"]:
                    c.fill = red_fill
                    c.font = Font(color=LEVEL_FONT["严重"])
        for col, width in zip("ABCDE", [10, 22, 16, 12, 30]):
            ws3.column_dimensions[col].width = width
        ws3.freeze_panes = "A2"

    # ---------- 4. 值比对（需确认） ----------
    if want_values:
        ws4 = wb.create_sheet("值比对(需确认)")
        for j, h in enumerate(["位号", "BOM 声明值", "原理图标注值", "是否一致", "说明"], 1):
            c = ws4.cell(1, j, h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=HEAD)
            c.border = border
        for i, d in enumerate(result.value_diffs, 2):
            values = [d["refs"], d["bom"], d["sch"], "一致" if d["equivalent"] else "★不一致", d["note"]]
            for j, v in enumerate(values, 1):
                c = ws4.cell(i, j, v)
                c.border = border
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if not d["equivalent"]:
                    c.fill, c.font = yel_fill, yel_font
        for col, width in zip("ABCDE", [16, 24, 18, 10, 40]):
            ws4.column_dimensions[col].width = width
        ws4.freeze_panes = "A2"

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)
    return out_path
